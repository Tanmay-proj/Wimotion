from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation(r'E:\WIFI_INTERFEROMETRIC_NEURAL_SENSING\intermediate_ppt.pptx')
print(f"Total Slides: {len(prs.slides)}\n")

for idx, s in enumerate(prs.slides, 1):
    print(f"==================================================")
    print(f"SLIDE {idx} (Shapes count: {len(s.shapes)})")
    print(f"==================================================")
    for i, sh in enumerate(s.shapes):
        left = sh.left / 914400 if sh.left is not None else 0
        top = sh.top / 914400 if sh.top is not None else 0
        width = sh.width / 914400 if sh.width is not None else 0
        height = sh.height / 914400 if sh.height is not None else 0
        
        # Color & Fill info
        fill_type = sh.fill.type if hasattr(sh, 'fill') else None
        
        print(f"\n--- Shape {i:02d}: Name='{sh.name}', Type={sh.shape_type} ---")
        print(f"    Position: left={left:.3f}in, top={top:.3f}in, w={width:.3f}in, h={height:.3f}in")
        
        if sh.has_text_frame:
            tf = sh.text_frame
            print(f"    TextFrame Margins: L={tf.margin_left/914400 if tf.margin_left else 0:.3f}in, R={tf.margin_right/914400 if tf.margin_right else 0:.3f}in, T={tf.margin_top/914400 if tf.margin_top else 0:.3f}in, B={tf.margin_bottom/914400 if tf.margin_bottom else 0:.3f}in, wrap={tf.word_wrap}")
            for p_idx, p in enumerate(tf.paragraphs):
                p_text = p.text.strip().encode('ascii', 'ignore').decode('ascii')
                sz = p.font.size.pt if (p.font and p.font.size) else 'inherit'
                fn = p.font.name if (p.font and p.font.name) else 'inherit'
                b = p.font.bold if p.font else 'inherit'
                c = p.font.color.rgb if (p.font and p.font.color and p.font.color.type == 1) else 'inherit'
                align = p.alignment
                print(f"      P{p_idx} [sz={sz}, fn={fn}, b={b}, col={c}, align={align}]: \"{p_text}\"")
                for r_idx, r in enumerate(p.runs):
                    r_text = r.text.encode('ascii', 'ignore').decode('ascii')
                    r_sz = r.font.size.pt if (r.font and r.font.size) else 'inherit'
                    r_fn = r.font.name if (r.font and r.font.name) else 'inherit'
                    r_b = r.font.bold if r.font else 'inherit'
                    r_c = r.font.color.rgb if (r.font and r.font.color and r.font.color.type == 1) else 'inherit'
                    if r_text.strip():
                        print(f"          R{r_idx} [sz={r_sz}, fn={r_fn}, b={r_b}, col={r_c}]: \"{r_text}\"")
