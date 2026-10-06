from pptx import Presentation

prs = Presentation(r'E:\WIFI_INTERFEROMETRIC_NEURAL_SENSING\intermediate_ppt.pptx')
for idx in [9, 10]:
    s = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for i, sh in enumerate(s.shapes):
        t = sh.text_frame.text[:30].strip().replace('\n', ' ') if sh.has_text_frame else ''
        t_clean = t.encode('ascii', 'ignore').decode('ascii')
        print(f"sh {i}: type={sh.shape_type}, left={sh.left/914400:.2f}in, top={sh.top/914400:.2f}in, w={sh.width/914400:.2f}in, txt='{t_clean}'")
